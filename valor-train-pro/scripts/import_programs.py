#!/usr/bin/env python3
"""
Import Corey's training programs from his tracking workbook into Valor Train Pro.

Reads the "Program - ..." tabs (not the roster or logs) and writes program_templates rows:
  In-season: "Program - Monday" + "Program - Wednesday"  (12 weeks, blocks of 4)
  Off-season: "Program - Day 1 (Mon)" / "Day 2 (Tue-Wed)" / "Day 3 (Thu)"  (weeks 1-4)
Each exercise keeps its group (primary / superset / finisher / extra), the prescription text,
which weeks it runs, and the target written for each week ("3 x 55%", "x10", "Bodyweight"...).

  pip install openpyxl   (a venv is fine)
  set -a; source ~/Documents/Claude/Projects/playbook/.env; set +a
  python3 scripts/import_programs.py "<workbook.xlsx>" [--dry-run]
Re-running replaces templates that have the same name (assignments on them are kept only if unchanged names).
"""
import json, os, re, sys, urllib.request
import openpyxl

URL = "https://gpotwyuttkkygvxzktep.supabase.co"
WEEKDAY = {"mon": 1, "tue": 2, "wed": 3, "thu": 4, "fri": 5, "sat": 6, "sun": 7}

def clean(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v not in (None, "") else ""

def group_of(prescription):
    p = prescription.lower()
    for g in ("primary", "superset", "finisher", "extra"):
        if p.startswith(g):
            return g
    return "other"

def parse_day(ws, key, label, weekday, max_week):
    header_row = None
    for r in range(1, 12):
        if clean(ws.cell(r, 1).value).lower() == "exercise":
            header_row = r; break
    week_cols = {}
    for c in range(3, ws.max_column + 1):
        m = re.match(r"wk\s*(\d+)", clean(ws.cell(header_row, c).value).lower())
        if m: week_cols[int(m.group(1))] = c
    day = {"key": key, "label": label, "weekday": weekday, "warmup": clean(ws["A2"].value), "cue": "", "finish": "", "exercises": []}
    block_weeks, block_label = list(range(1, max_week + 1)), ""
    for r in range(header_row + 1, ws.max_row + 1):
        a = clean(ws.cell(r, 1).value)
        if not a:
            continue
        up = a.upper()
        m = re.search(r"WEEKS?\s*(\d+)\s*-\s*(\d+)", up)
        if m and ("BLOCK" in up or "·" in a or up.startswith("WEEKS")):
            block_weeks = list(range(int(m.group(1)), int(m.group(2)) + 1))
            fm = re.search(r"\(([^)]+)\)", a) or re.search(r"·\s*(.+)$", a)
            block_label = fm.group(1).strip().title() if fm else ""
            continue
        if up.startswith("COACH'S CUE"):
            day["cue"] = a.split(":", 1)[1].strip().strip('"') if ":" in a else a; continue
        if a.lower().startswith("finish with"):
            day["finish"] = a; continue
        prescription = clean(ws.cell(r, 2).value)
        targets = {str(w): clean(ws.cell(r, c).value) for w, c in week_cols.items() if w in block_weeks and clean(ws.cell(r, c).value)}
        day["exercises"].append({"name": a, "group": group_of(prescription), "prescription": prescription,
                                 "block": block_label, "weeks": block_weeks, "targets": targets})
    return day

def build(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    names = wb.sheetnames
    out = []
    if "Program - Monday" in names and "Program - Wednesday" in names:
        out.append({"name": "Fall In-Season Strength (Mon/Wed, 12 weeks)", "season": "in_season", "weeks": 12,
                    "description": "In-Season I, Phase 3 (Strength/Power). Three 4-week blocks; Monday and Wednesday alternate Total Body and Lower Body. Percentages are of each athlete's most recent tested or estimated max; if bar speed slows, hold the load.",
                    "days": [parse_day(wb["Program - Monday"], "mon", "Monday", 1, 12), parse_day(wb["Program - Wednesday"], "wed", "Wednesday", 3, 12)]})
    off = [n for n in names if n.startswith("Program - Day ")]
    if off:
        days = []
        for n in sorted(off):
            m = re.match(r"Program - Day (\d) \(([^)]+)\)", n)
            num, when = m.group(1), m.group(2)
            wd = WEEKDAY.get(when.split("-")[0][:3].lower(), 1)
            days.append(parse_day(wb[n], f"day{num}", f"Day {num} ({when})", wd, 4))
        out.append({"name": "Off-Season Foundation (3 days, weeks 1-4)", "season": "off_season", "weeks": 4,
                    "description": "Off-season Foundation block. Day 1 Monday (Total Body), Day 2 Tuesday or Wednesday, athlete's choice (Lower Body), Day 3 Thursday (Upper Body). Start each session with the mobility warm-up and SAQ sheet.",
                    "days": days})
    return out

def main():
    path = sys.argv[1]; dry = "--dry-run" in sys.argv
    templates = build(path)
    for t in templates:
        n = sum(len(d["exercises"]) for d in t["days"])
        print(f"{t['name']}: {len(t['days'])} days, {n} exercise rows, {t['weeks']} weeks")
        for d in t["days"]:
            print(f"   {d['label']}: {len(d['exercises'])} exercises, blocks {sorted({e['block'] for e in d['exercises']})}")
    if dry:
        print(json.dumps(templates[0]["days"][0]["exercises"][:3], indent=1)); return
    key = os.environ["VALOR_TRAINPRO_SUPABASE_SERVICE_ROLE_KEY"]
    h = {"apikey": key, "Authorization": f"Bearer {key}", "content-type": "application/json", "Prefer": "return=representation"}
    for t in templates:
        t["source"] = os.path.basename(path)
        q = urllib.request.quote(t["name"])
        existing = json.loads(urllib.request.urlopen(urllib.request.Request(f"{URL}/rest/v1/program_templates?select=id&name=eq.{q}", headers=h)).read())
        if existing:
            req = urllib.request.Request(f"{URL}/rest/v1/program_templates?id=eq.{existing[0]['id']}", data=json.dumps(t).encode(), headers=h, method="PATCH")
        else:
            req = urllib.request.Request(f"{URL}/rest/v1/program_templates", data=json.dumps(t).encode(), headers=h, method="POST")
        row = json.loads(urllib.request.urlopen(req).read())[0]
        print(("updated " if existing else "created ") + row["id"], t["name"])

if __name__ == "__main__":
    main()
