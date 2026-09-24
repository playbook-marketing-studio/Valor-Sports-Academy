#!/usr/bin/env python3
"""
Import Corey's current athletes from the workbook's Roster tab into Valor Train Pro.
DRY RUN by default: prints what it would create. Add --apply to write.

Fields: name, sport, season (which roster section), frequency, class days, nutrition plan.
The workbook has no parent contact or age, so athletes land without a parent login; add the
parent email on each athlete's Profile (Edit) and send a login from the app.
Skips anyone whose first + last name already exists.

  set -a; source ~/Documents/Claude/Projects/playbook/.env; set +a
  python3 scripts/import_roster.py "<workbook.xlsx>"            # preview
  python3 scripts/import_roster.py "<workbook.xlsx>" --apply    # write
"""
import json, os, re, sys, urllib.request
import openpyxl

URL = "https://gpotwyuttkkygvxzktep.supabase.co"
clean = lambda v: re.sub(r"\s+", " ", str(v)).strip() if v not in (None, "") else ""

def read(path):
    ws = openpyxl.load_workbook(path, data_only=True)["Roster"]
    out, season = [], None
    for r in range(1, ws.max_row + 1):
        a = clean(ws.cell(r, 1).value)
        if a.upper().startswith("IN-SEASON ROSTER"): season = "in_season"; continue
        if a.upper().startswith("OFF-SEASON ROSTER"): season = "off_season"; continue
        if not a or a.lower() == "athlete name" or not season: continue
        first, *rest = a.split(" ")
        out.append({
            "first_name": first, "last_name": " ".join(rest) or None,
            "sport": clean(ws.cell(r, 2).value) or None, "season": season,
            "frequency": clean(ws.cell(r, 3).value) or None, "class_days": clean(ws.cell(r, 4).value) or None,
            "nutrition_plan": clean(ws.cell(r, 5).value).lower() == "yes",
        })
    return out

def main():
    rows = read(sys.argv[1]); apply = "--apply" in sys.argv
    key = os.environ["VALOR_TRAINPRO_SUPABASE_SERVICE_ROLE_KEY"]
    h = {"apikey": key, "Authorization": f"Bearer {key}", "content-type": "application/json", "Prefer": "return=representation"}
    existing = json.loads(urllib.request.urlopen(urllib.request.Request(f"{URL}/rest/v1/athletes?select=first_name,last_name", headers=h)).read())
    have = {(e["first_name"].lower(), (e["last_name"] or "").lower()) for e in existing}
    new = [r for r in rows if (r["first_name"].lower(), (r["last_name"] or "").lower()) not in have]
    print(f"{len(rows)} on the roster, {len(rows) - len(new)} already in the app, {len(new)} to add")
    for r in new:
        print(f"  + {r['first_name']} {r['last_name'] or ''} · {r['sport'] or '-'} · {r['season']} · {r['frequency'] or '-'} · {r['class_days'] or '-'} · nutrition {'yes' if r['nutrition_plan'] else 'no'}")
    if not apply:
        print("dry run: nothing written. Re-run with --apply."); return
    if new:
        urllib.request.urlopen(urllib.request.Request(f"{URL}/rest/v1/athletes", data=json.dumps(new).encode(), headers=h, method="POST"))
        print(f"added {len(new)}")

if __name__ == "__main__":
    main()
